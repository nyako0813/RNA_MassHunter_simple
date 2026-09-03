from pathlib import Path

import openpyxl
import pytest

from rna_masshunter.config import load_config
from rna_masshunter.excel_report import write_simple_mass_hunter_report
from rna_masshunter.mass_comparison import MassComparisonRow
from rna_masshunter.models import Fragment

REPO_ROOT = Path(__file__).resolve().parent.parent


def _row(fragment_id, charge, delta_da, recommended_formula=None, known_modification=None) -> MassComparisonRow:
    return MassComparisonRow(
        peak_id="PK0001", fragment_id=fragment_id, sequence="GG", charge=charge, intensity=1000.0,
        observed_mz=500.0, observed_mass=1000.0, theoretical_mass=1000.0 - delta_da,
        delta_da=delta_da, delta_ppm=0.0,
        recommended_formula=recommended_formula, formula_candidates="",
        known_modification=known_modification, modification_candidates="",
    )


def _fragment() -> Fragment:
    return Fragment(
        fragment_id="F1", target_id="T", sequence="GG", start=1, end=2, standard_start=1, standard_end=2,
        enzyme="RNase_T1", missed_cleavages=0, terminal_form="default", unmodified_mass=692.0,
    )


def test_visualization_sources_from_mass_comparison_and_color_codes_by_candidate(tmp_path):
    rows = [
        _row("F1", 2, 0.0, recommended_formula="0"),
        _row("F1", 2, 14.0156, known_modification="1-methyladenosine"),
        _row("F1", 3, 0.5),  # no candidate
    ]
    config = load_config(REPO_ROOT / "config.yaml")

    out = write_simple_mass_hunter_report(tmp_path / "report.xlsx", config, [_fragment()], [], rows, [])

    wb = openpyxl.load_workbook(out)
    ws8 = wb["08_Visualization"]
    assert len(ws8._charts) == 1
    assert len(ws8._charts[0].series) == 2

    header = [ws8.cell(row=1, column=c).value for c in range(20, 24)]
    assert header == ["Charge (has candidate)", "ΔDa (has candidate)", "Charge (no candidate)", "ΔDa (no candidate)"]

    has_candidate_deltas = {ws8.cell(row=r, column=21).value for r in (2, 3) if ws8.cell(row=r, column=21).value is not None}
    assert has_candidate_deltas == {0.0, 14.0156}
    no_candidate_deltas = {ws8.cell(row=r, column=23).value for r in (2, 3) if ws8.cell(row=r, column=23).value is not None}
    assert no_candidate_deltas == {0.5}


def test_visualization_disabled_writes_placeholder_only(tmp_path):
    rows = [_row("F1", 2, 0.0, recommended_formula="0")]
    config = load_config(REPO_ROOT / "config.yaml")
    config.visualization = {"enabled": False}

    out = write_simple_mass_hunter_report(tmp_path / "report.xlsx", config, [_fragment()], [], rows, [])

    wb = openpyxl.load_workbook(out)
    ws8 = wb["08_Visualization"]
    assert len(ws8._charts) == 0
    assert ws8["A3"].value == "Visualization is disabled (config.visualization.enabled = false)."


def test_visualization_no_rows_writes_placeholder_only(tmp_path):
    config = load_config(REPO_ROOT / "config.yaml")

    out = write_simple_mass_hunter_report(tmp_path / "report.xlsx", config, [_fragment()], [], [], [])

    wb = openpyxl.load_workbook(out)
    ws8 = wb["08_Visualization"]
    assert len(ws8._charts) == 0
    assert ws8["A3"].value == "No Mass Comparison rows to plot."
