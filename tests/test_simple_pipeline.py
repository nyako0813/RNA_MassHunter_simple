"""End-to-end integration test for `rna_masshunter.simple_pipeline.run()`.

Builds a minimal synthetic mzML by hand (pyteomics has no mzML writer, so
this constructs the XML directly: base64/float64-encoded binary arrays plus
just enough cvParam elements for `pyteomics.mzml.MzML` — which itself
requires the `psims` package, see requirements.txt — to parse ms level, m/z
array, intensity array, scan start time, and (for MS2) precursor m/z/charge).
Everything the test asserts against (fragment mass, theoretical MS2 ion m/z)
is computed via the real vendored/new functions rather than hardcoded, so
the test stays correct if e.g. base_masses.yaml or the digestion rules
change.
"""
from __future__ import annotations

import base64
import struct
import textwrap
from pathlib import Path

import openpyxl
import pytest

from rna_masshunter import simple_pipeline
from rna_masshunter.config import load_config, resolve_paths
from rna_masshunter.digestion import digest_sequence
from rna_masshunter.masses import load_base_masses, mz_from_neutral_mass
from rna_masshunter.ms2_extraction import generate_theoretical_ms2_ions

REPO_ROOT = Path(__file__).resolve().parent.parent
SEQUENCE = "GGGCGUGUGGCGUAGUCGGUAGCGCGCUCCCUUAGCAUGGGAGAGGUCUCCGGUUCGAUUCCGGACUCGUCCACCA"
M1A_SHIFT_DA = 14.0156  # matches data/modifications.yaml's m1A mass_shift_from_unmodified


# --- minimal synthetic mzML writer (pyteomics has no writer) ----------------

def _encode_f64(values: list[float]) -> str:
    return base64.b64encode(struct.pack(f"<{len(values)}d", *values)).decode("ascii")


def _binary_array_xml(values: list[float], accession: str, name: str) -> str:
    encoded = _encode_f64(values)
    return (
        f'<binaryDataArray encodedLength="{len(encoded)}">'
        '<cvParam cvRef="MS" accession="MS:1000523" name="64-bit float" value=""/>'
        '<cvParam cvRef="MS" accession="MS:1000576" name="no compression" value=""/>'
        f'<cvParam cvRef="MS" accession="{accession}" name="{name}" value=""/>'
        f"<binary>{encoded}</binary></binaryDataArray>"
    )


def _spectrum_xml(index, scan_id, ms_level, rt_minutes, mzs, intensities, precursor_mz=None, precursor_charge=None) -> str:
    precursor_xml = ""
    if precursor_mz is not None:
        precursor_xml = (
            '<precursorList count="1"><precursor><selectedIonList count="1"><selectedIon>'
            f'<cvParam cvRef="MS" accession="MS:1000744" name="selected ion m/z" value="{precursor_mz}"/>'
            f'<cvParam cvRef="MS" accession="MS:1000041" name="charge state" value="{precursor_charge}"/>'
            "</selectedIon></selectedIonList></precursor></precursorList>"
        )
    return (
        f'<spectrum index="{index}" id="{scan_id}" defaultArrayLength="{len(mzs)}">'
        f'<cvParam cvRef="MS" accession="MS:1000511" name="ms level" value="{ms_level}"/>'
        '<scanList count="1"><cvParam cvRef="MS" accession="MS:1000795" name="no combination" value=""/>'
        '<scan><cvParam cvRef="MS" accession="MS:1000016" name="scan start time" '
        f'value="{rt_minutes}" unitCvRef="UO" unitAccession="UO:0000031" unitName="minute"/></scan></scanList>'
        f"{precursor_xml}"
        '<binaryDataArrayList count="2">'
        f'{_binary_array_xml(mzs, "MS:1000514", "m/z array")}'
        f'{_binary_array_xml(intensities, "MS:1000515", "intensity array")}'
        "</binaryDataArrayList></spectrum>"
    )


def _write_mzml(path: Path, spectra: list[dict]) -> None:
    body = "".join(
        _spectrum_xml(i, s["id"], s["ms_level"], s["rt"], s["mzs"], s["intensities"], s.get("precursor_mz"), s.get("precursor_charge"))
        for i, s in enumerate(spectra)
    )
    xml = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<mzML xmlns="http://psi.hupo.org/ms/mzml" version="1.1.0">'
        '<cvList count="2">'
        '<cv id="MS" fullName="PSI-MS" version="4.1.0" URI="https://raw.githubusercontent.com/HUPO-PSI/psi-ms-CV/master/psi-ms.obo"/>'
        '<cv id="UO" fullName="UNIT-ONTOLOGY" version="09:04:2014" URI="http://obo.cvs.sourceforge.net/obo/obo/ontology/phenotype/unit.obo"/>'
        "</cvList>"
        '<fileDescription><fileContent><cvParam cvRef="MS" accession="MS:1000579" name="MS1 spectrum" value=""/></fileContent></fileDescription>'
        '<softwareList count="1"><software id="sw1" version="1.0"/></softwareList>'
        '<instrumentConfigurationList count="1"><instrumentConfiguration id="IC1"/></instrumentConfigurationList>'
        '<dataProcessingList count="1"><dataProcessing id="dp1"><processingMethod order="1" softwareRef="sw1"/></dataProcessing></dataProcessingList>'
        f'<run id="run1"><spectrumList count="{len(spectra)}" defaultDataProcessingRef="dp1">{body}</spectrumList></run>'
        "</mzML>"
    )
    path.write_text(xml, encoding="utf-8")


# --- fixture: figure out the target fragment + its theoretical MS2 ions -----

@pytest.fixture(scope="module")
def target_fragment_and_ions():
    """Digest SEQUENCE with the same settings the test config.yaml uses, and
    compute the first fragment's theoretical d/w/a/z ions for real (not
    hardcoded), so the synthetic MS2 peaks below are guaranteed to match."""
    config = load_config(REPO_ROOT / "config.yaml")
    config.sequence["sequence"] = SEQUENCE
    config.digestion.update({"enzyme": "RNase_T1", "missed_cleavages": 1, "min_length": 2})
    base_masses = load_base_masses(REPO_ROOT / "data" / "base_masses.yaml")

    fragments = digest_sequence(
        target_id="IntegrationTest", sequence=SEQUENCE,
        position_map={i: i for i in range(1, len(SEQUENCE) + 1)},
        config=config, base_masses=base_masses,
    )
    fragment = fragments[0]
    assert len(fragment.sequence) >= 2  # otherwise generate_theoretical_ms2_ions would yield nothing

    ions = generate_theoretical_ms2_ions([fragment], config, base_masses)
    assert ions, "expected at least one theoretical MS2 ion for the target fragment"
    return fragment, ions


@pytest.fixture
def config_yaml_text(tmp_path):
    def _make(mzml_path: Path, output_dir: Path, ms2_enabled: bool) -> Path:
        text = textwrap.dedent(f"""
            sequence:
              name: IntegrationTest
              sequence: {SEQUENCE}
            instrument:
              polarity: negative
            digestion:
              enzyme: RNase_T1
              missed_cleavages: 1
              min_length: 2
            input:
              mzml_path: {mzml_path}
            project:
              output_dir: {output_dir}
            ms1_peak_extraction:
              mz_min: 0
              mz_max: 5000
              intensity_threshold: 0
            ms2_annotation:
              enabled: {"true" if ms2_enabled else "false"}
            reporting:
              excel_output: true
        """)
        config_path = tmp_path / f"config_{'ms2' if ms2_enabled else 'ms1only'}.yaml"
        config_path.write_text(text, encoding="utf-8")
        return config_path

    return _make


def _build_synthetic_mzml(tmp_path, fragment, ions) -> Path:
    unmodified_mz = mz_from_neutral_mass(fragment.unmodified_mass, 2, "negative")
    modified_mz = mz_from_neutral_mass(fragment.unmodified_mass + M1A_SHIFT_DA, 2, "negative")

    ms2_ions = [ion for ion in ions if ion.parent_fragment_id == fragment.fragment_id][:2]
    ms2_mzs = [ion.theoretical_mz for ion in ms2_ions]

    mzml_path = tmp_path / "synthetic.mzML"
    _write_mzml(mzml_path, [
        {"id": "scan=1", "ms_level": 1, "rt": 3.0, "mzs": [unmodified_mz], "intensities": [1000.0]},
        {"id": "scan=2", "ms_level": 1, "rt": 3.5, "mzs": [modified_mz], "intensities": [1200.0]},
        {
            "id": "scan=3", "ms_level": 2, "rt": 3.5, "mzs": ms2_mzs, "intensities": [500.0] * len(ms2_mzs),
            "precursor_mz": modified_mz, "precursor_charge": 2,
        },
    ])
    return mzml_path


def test_pipeline_runs_end_to_end_with_ms2(tmp_path, config_yaml_text, target_fragment_and_ions):
    fragment, ions = target_fragment_and_ions
    output_dir = tmp_path / "output"
    mzml_path = _build_synthetic_mzml(tmp_path, fragment, ions)
    config_path = config_yaml_text(mzml_path, output_dir, ms2_enabled=True)

    result = simple_pipeline.run(config_path, project_root=REPO_ROOT)

    assert result["output_path"]
    output_path = Path(result["output_path"])
    assert output_path.exists()

    rows = result["mass_comparison_rows"]
    assert rows
    matching_rows = [row for row in rows if row.fragment_id == fragment.fragment_id and row.charge == 2]
    assert matching_rows

    unmodified_row = next((row for row in matching_rows if row.delta_da == pytest.approx(0.0, abs=1e-3)), None)
    assert unmodified_row is not None
    assert unmodified_row.recommended_formula == "0"
    # Not asserting known_modification is None here: pseudouridine (an
    # isomer, mass_shift_from_unmodified == 0) is a legitimate Modification
    # Candidate match at ΔDa == 0 — that's correct behavior of the known-
    # modification search, not a bug.

    modified_row = next((row for row in matching_rows if row.delta_da == pytest.approx(M1A_SHIFT_DA, abs=1e-2)), None)
    assert modified_row is not None
    assert modified_row.known_modification is not None
    assert "m1A" in modified_row.modification_candidates
    # MS2 reference info should have attached to the modified-precursor row
    # (that's the one the synthetic MS2 spectrum's precursor_mz matches),
    # and it must not have changed the Formula/Modification verdict (原則3).
    assert modified_row.ms2_matched_ion_count > 0
    assert modified_row.ms2_matched_ions != ""
    assert "m1A" in modified_row.modification_candidates

    wb = openpyxl.load_workbook(output_path)
    assert wb.sheetnames == ["01_Index", "02_Input", "03_Theoretical", "04_Observed_Mass", "05_Mass_Intensity", "06_Mass_Comparison", "07_Modifications", "08_Visualization"]
    ms2_header = [cell.value for cell in wb["06_Mass_Comparison"][3]]
    assert "MS2 Spectrum ID" in ms2_header
    assert "MS2 Matched Ions" in ms2_header
    # 01_Index backlinks/hyperlinks are wired (§16.3)
    assert wb["02_Input"]["A1"].value == "← Back to Index"
    assert wb["02_Input"]["A1"].hyperlink is not None


def test_pipeline_ms2_disabled_matches_ms1_only_behavior(tmp_path, config_yaml_text, target_fragment_and_ions):
    """仕様書 §20: config.ms2_annotation.enabled=false でもMS1のみのフロー
    と同じ結果になり（回帰しない）、MS2列は常に既定値のまま。"""
    fragment, ions = target_fragment_and_ions
    output_dir = tmp_path / "output"
    mzml_path = _build_synthetic_mzml(tmp_path, fragment, ions)
    config_path = config_yaml_text(mzml_path, output_dir, ms2_enabled=False)

    result = simple_pipeline.run(config_path, project_root=REPO_ROOT)

    assert result["ms2_spectra"] == []
    rows = result["mass_comparison_rows"]
    assert rows
    for row in rows:
        assert row.ms2_spectrum_id is None
        assert row.ms2_matched_ion_count == 0
        assert row.ms2_matched_ions == ""

    modified_row = next(
        (row for row in rows if row.fragment_id == fragment.fragment_id and row.charge == 2 and row.delta_da == pytest.approx(M1A_SHIFT_DA, abs=1e-2)),
        None,
    )
    assert modified_row is not None
    assert modified_row.known_modification is not None
    assert "m1A" in modified_row.modification_candidates


def test_pipeline_dry_run_without_mzml_completes_with_warnings(tmp_path, config_yaml_text):
    output_dir = tmp_path / "output"
    config_path = config_yaml_text(mzml_path="", output_dir=output_dir, ms2_enabled=True)

    result = simple_pipeline.run(config_path, project_root=REPO_ROOT)

    assert result["fragments"]
    assert result["peaks"] == []
    assert result["mass_comparison_rows"] == []
    assert any(w["Source"] == "simple_pipeline" and "mzml_path" in w["Message"] for w in result["warnings"])
    # Excel is still produced (with 04-06 sheets empty) even without mzML.
    assert result["output_path"]
    assert Path(result["output_path"]).exists()


def test_pipeline_merge_profile_points_toggle(tmp_path):
    """仕様書 §14B: config.ms1_peak_extraction.merge_profile_points が
    true（既定）なら1スキャン内の隣接プロファイル点が統合され、false なら
    生のまま全点が残ることを、simple_pipeline.run() を通しで確認する。
    §14Cのmerge_across_scansはscan-rank差0（同一scan）も対象に含むため
    （境界を跨がない場合の当然の帰結）、ここでは両方を明示的にfalseに
    しないとmerge_profile_points単体の効果が見えないので、falseケースは
    両方offにする。"""
    mzml_path = tmp_path / "profile.mzML"
    _write_mzml(mzml_path, [
        {
            "id": "scan=1", "ms_level": 1, "rt": 2.0,
            # a 5-point profile-mode "mountain" around one true ion, plus one
            # well-separated second ion in the same scan.
            "mzs": [600.00000, 600.00002, 600.00004, 600.00006, 600.00008, 900.0],
            "intensities": [100.0, 400.0, 900.0, 350.0, 90.0, 700.0],
        },
    ])

    def _config_path(merge_enabled: bool) -> Path:
        text = textwrap.dedent(f"""
            sequence:
              name: MergeToggleTest
              sequence: {SEQUENCE}
            instrument:
              polarity: negative
            input:
              mzml_path: {mzml_path}
            project:
              output_dir: {tmp_path / ('merge_on' if merge_enabled else 'merge_off')}
            ms1_peak_extraction:
              mz_min: 0
              mz_max: 5000
              intensity_threshold: 0
              merge_profile_points: {"true" if merge_enabled else "false"}
              merge_tolerance_ppm: 10
              merge_across_scans: {"true" if merge_enabled else "false"}
            ms2_annotation:
              enabled: false
            reporting:
              excel_output: false
        """)
        path = tmp_path / f"config_merge_{merge_enabled}.yaml"
        path.write_text(text, encoding="utf-8")
        return path

    merged_result = simple_pipeline.run(_config_path(True), project_root=REPO_ROOT)
    unmerged_result = simple_pipeline.run(_config_path(False), project_root=REPO_ROOT)

    assert len(unmerged_result["peaks"]) == 6
    assert len(merged_result["peaks"]) == 2
    merged_intensities = sorted(p.intensity for p in merged_result["peaks"])
    assert merged_intensities == [700.0, 900.0]


def test_pipeline_merge_across_scans_collapses_elution_profile_into_one_row(tmp_path, target_fragment_and_ions):
    """仕様書 §14C の統合テスト: 実データ（19 Gln2h.mzML）で確認されたのと
    同じ状況——同一イオンが複数の連続scanにまたがって検出される——を合成
    mzMLで再現し、06_Mass_Comparisonが1行に統合され、Scan Count/RT Range
    列にその情報が残ることを確認する。"""
    fragment, _ions = target_fragment_and_ions
    unmodified_mz = mz_from_neutral_mass(fragment.unmodified_mass, 2, "negative")

    mzml_path = tmp_path / "elution.mzML"
    scan_rts = [3.00, 3.02, 3.04, 3.06, 3.08]
    intensities = [200.0, 600.0, 1000.0, 550.0, 180.0]  # apex mid-elution
    _write_mzml(mzml_path, [
        {"id": f"scan={i + 1}", "ms_level": 1, "rt": rt, "mzs": [unmodified_mz], "intensities": [intensity]}
        for i, (rt, intensity) in enumerate(zip(scan_rts, intensities))
    ])

    config_text = textwrap.dedent(f"""
        sequence:
          name: IntegrationTest
          sequence: {SEQUENCE}
        instrument:
          polarity: negative
        digestion:
          enzyme: RNase_T1
          missed_cleavages: 1
          min_length: 2
        input:
          mzml_path: {mzml_path}
        project:
          output_dir: {tmp_path / "output"}
        ms1_peak_extraction:
          mz_min: 0
          mz_max: 5000
          intensity_threshold: 0
          merge_max_scan_gap: 2
        ms2_annotation:
          enabled: false
        reporting:
          excel_output: true
    """)
    config_path = tmp_path / "config_elution.yaml"
    config_path.write_text(config_text, encoding="utf-8")

    result = simple_pipeline.run(config_path, project_root=REPO_ROOT)

    matching_rows = [row for row in result["mass_comparison_rows"] if row.fragment_id == fragment.fragment_id and row.charge == 2]
    assert len(matching_rows) == 1
    row = matching_rows[0]
    assert row.scan_count == 5
    assert row.rt_range == pytest.approx((3.00, 3.08))
    assert row.intensity == 1000.0  # the apex scan was kept as representative

    wb = openpyxl.load_workbook(result["output_path"])
    ws6 = wb["06_Mass_Comparison"]
    header = [cell.value for cell in ws6[3]]
    assert "Scan Count" in header
    assert "RT Range" in header
    scan_count_col = header.index("Scan Count")
    rt_range_col = header.index("RT Range")
    data_row = next(r for r in ws6.iter_rows(min_row=4, values_only=True) if r[1] == fragment.fragment_id and r[3] == 2)
    assert data_row[scan_count_col] == 5
    assert data_row[rt_range_col] == "3.000-3.080"


# --- §24 (Phase 10): Nuclease P1 complete-digestion mode --------------------

def _p1_config_text(mzml_path, output_dir, ap_enabled: bool) -> str:
    return textwrap.dedent(f"""
        sequence:
          name: should_be_ignored_in_p1_mode
          sequence: {SEQUENCE}
        instrument:
          polarity: negative
        digestion:
          enzyme: Nuclease_P1
        alkaline_phosphatase:
          enabled: {"true" if ap_enabled else "false"}
        input:
          mzml_path: {mzml_path}
        project:
          output_dir: {output_dir}
        ms1_peak_extraction:
          mz_min: 0
          mz_max: 5000
          intensity_threshold: 0
        reporting:
          excel_output: true
    """)


def test_p1_mode_skips_sequence_cca_and_fragments(tmp_path):
    config_path = tmp_path / "config_p1_dry.yaml"
    config_path.write_text(_p1_config_text(mzml_path="", output_dir=tmp_path / "output", ap_enabled=True), encoding="utf-8")

    result = simple_pipeline.run(config_path, project_root=REPO_ROOT)

    assert result["mode"] == "p1_nucleoside"
    assert result["fragments"] == []
    assert result["cca_result"] is None
    assert result["theoretical_mass"] is None
    assert len(result["nucleoside_targets"]) > 100  # 4 standard + ~118 modified
    assert any(t.label == "A" and t.category == "standard" for t in result["nucleoside_targets"])


def test_p1_mode_matches_standard_and_modified_nucleosides_end_to_end(tmp_path):
    from rna_masshunter.masses import mz_from_neutral_mass
    from rna_masshunter.nucleoside_targets import build_nucleoside_target_universe
    from rna_masshunter.modifications import load_modifications

    modifications = load_modifications(REPO_ROOT / "data" / "modifications.yaml")
    targets = build_nucleoside_target_universe(modifications)
    adenosine = next(t for t in targets if t.label == "A")
    m1a = next(t for t in targets if t.label == "m1A")

    mzml_path = tmp_path / "p1.mzML"
    _write_mzml(mzml_path, [
        {"id": "scan=1", "ms_level": 1, "rt": 1.0, "mzs": [mz_from_neutral_mass(adenosine.nucleoside_mass, 1, "negative")], "intensities": [1000.0]},
        {"id": "scan=2", "ms_level": 1, "rt": 1.5, "mzs": [mz_from_neutral_mass(m1a.nucleoside_mass, 1, "negative")], "intensities": [800.0]},
    ])
    config_path = tmp_path / "config_p1.yaml"
    config_path.write_text(_p1_config_text(mzml_path, tmp_path / "output", ap_enabled=True), encoding="utf-8")

    result = simple_pipeline.run(config_path, project_root=REPO_ROOT)

    rows = result["nucleoside_comparison_rows"]
    standard_match = next((r for r in rows if r.target_label == "A"), None)
    modified_match = next((r for r in rows if r.target_label == "m1A"), None)
    assert standard_match is not None
    assert standard_match.target_category == "standard"
    assert standard_match.delta_da == pytest.approx(0.0, abs=1e-6)
    assert modified_match is not None
    assert modified_match.target_category == "modified"

    wb = openpyxl.load_workbook(result["output_path"])
    assert wb.sheetnames == ["01_Index", "02_Input", "03_Nucleoside_Targets", "04_Observed_Mass", "05_Mass_Intensity", "06_Nucleoside_Comparison", "07_Modifications", "08_Visualization"]
    ws6 = wb["06_Nucleoside_Comparison"]
    header = [cell.value for cell in ws6[3]]
    assert header == ["Peak ID", "Target Label", "Target Name", "Category", "Charge", "Intensity", "Observed Mass", "Theoretical Mass", "ΔDa", "Δppm", "Scan Count", "RT Range"]


def test_p1_mode_alkaline_phosphatase_disabled_matches_phosphorylated_mass(tmp_path):
    from rna_masshunter.masses import mz_from_neutral_mass, load_base_masses
    from rna_masshunter.nucleoside_targets import build_standard_nucleoside_targets

    base_masses = load_base_masses(REPO_ROOT / "data" / "base_masses.yaml")
    phosphate = float(base_masses["constants"]["phosphate"])
    adenosine = build_standard_nucleoside_targets()[0]
    assert adenosine.label == "A"

    mzml_path = tmp_path / "p1_phospho.mzML"
    _write_mzml(mzml_path, [
        {"id": "scan=1", "ms_level": 1, "rt": 1.0, "mzs": [mz_from_neutral_mass(adenosine.nucleoside_mass + phosphate, 1, "negative")], "intensities": [1000.0]},
    ])
    config_path = tmp_path / "config_p1_phospho.yaml"
    config_path.write_text(_p1_config_text(mzml_path, tmp_path / "output", ap_enabled=False), encoding="utf-8")

    result = simple_pipeline.run(config_path, project_root=REPO_ROOT)

    match = next((r for r in result["nucleoside_comparison_rows"] if r.target_label == "A"), None)
    assert match is not None
    assert match.delta_da == pytest.approx(0.0, abs=1e-6)
    assert match.theoretical_mass == pytest.approx(adenosine.nucleoside_mass + phosphate, abs=1e-6)
