"""Excel report writer for the simple RNA_MassHunter pipeline.

New file, NOT the same module as the original repository's much larger
`excel_report.py` (仕様書 §5 — a same-named-but-different file is
intentional: `write_excel_report()`'s 670-line body and its Intact/SCIEX/
shadow-audit sheet set target a different, unrelated output format). This
file only contains the new `write_simple_mass_hunter_report()` for the
01-08 sheet layout (仕様書 §16).

The handful of small formatting/index helpers below
(`_flatten_dict`, `_autosize_and_freeze`, `_sheet_link`, `_coerce_to_frame`,
`_json_safe`, `_excel_safe_cell`, `_add_index_and_backlinks`,
`_truncate_frame_if_needed`, `_append_excel_warning`, and the
`EXCEL_MAX_ROWS`/`DATA_START_ROW`/`EXCEL_DATA_ROW_LIMIT` constants) are
vendored from nyako0813/RNA_MassHunter `rna_masshunter/excel_report.py`
(仕様書 §4, §16.3 — reuse the existing Index/hyperlink/formatting
conventions rather than reimplementing them) — fetched read-only from the
original repo for reference, the same pattern already used for
models.py/config.py/ms2_extraction.py (see those files' docstrings). Only
these specific helpers were copied, not the surrounding 2314-line file or
any of the SCIEX/shadow-audit sheet-building logic.

One deliberate adaptation: the original hardcodes the index sheet's name as
"Index"; this project's sheet-naming convention numbers sheets ("01_Index",
"02_Input", ...), so `_add_index_and_backlinks` / `_autosize_and_freeze`
below take an `index_sheet_name` parameter instead of the literal string.
"""
from __future__ import annotations

import json
import math
from datetime import datetime
from enum import Enum
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter

from rna_masshunter.hypothesis_check import HypothesisCheckRow
from rna_masshunter.mass_comparison import MassComparisonRow, has_any_candidate
from rna_masshunter.models import Fragment, Modification, Peak, RunConfig
from rna_masshunter.nucleoside_comparison import NucleosideComparisonRow
from rna_masshunter.nucleoside_targets import NucleosideTarget, theoretical_mass_for_target
from rna_masshunter.observed_mass import assign_peak_ids

# --- vendored helpers (see module docstring) ---------------------------------

EXCEL_MAX_ROWS = 1_048_576
DATA_START_ROW = 3
EXCEL_DATA_ROW_LIMIT = EXCEL_MAX_ROWS - DATA_START_ROW


def _flatten_dict(data: dict[str, Any], prefix: str = "") -> list[dict[str, Any]]:
    rows = []
    for key, value in data.items():
        full_key = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            rows.extend(_flatten_dict(value, full_key))
        else:
            rows.append({"Parameter": full_key, "Value": value})
    return rows


def _autosize_and_freeze(writer: pd.ExcelWriter, index_sheet_name: str = "01_Index") -> None:
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2" if worksheet.title == index_sheet_name else "A4"
        for column_cells in worksheet.columns:
            max_length = 0
            column = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, min(len(value), 60))
            worksheet.column_dimensions[column].width = max(10, max_length + 2)


def _sheet_link(sheet_name: str, cell: str = "A1") -> str:
    return f"#'{sheet_name}'!{cell}"


def _coerce_to_frame(value: Any) -> pd.DataFrame:
    if isinstance(value, pd.DataFrame):
        return value
    if isinstance(value, list):
        return pd.DataFrame(value)
    if isinstance(value, dict):
        return pd.DataFrame([value])
    return pd.DataFrame([{"Value": value}])


def _json_safe(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, Enum):
        return _json_safe(value.value)
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if hasattr(value, "to_dict") and callable(value.to_dict):
        return _json_safe(value.to_dict())
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in sorted(value.items(), key=lambda pair: str(pair[0]))}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if isinstance(value, (set, frozenset)):
        normalized = [_json_safe(item) for item in value]
        return sorted(normalized, key=lambda item: json.dumps(item, ensure_ascii=False, sort_keys=True, default=str))
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _excel_safe_cell(value: Any) -> Any:
    normalized = _json_safe(value)
    if normalized is None:
        return ""
    if isinstance(normalized, (list, dict)):
        return json.dumps(normalized, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return normalized


def _add_index_and_backlinks(writer: pd.ExcelWriter, sheet_names: list[str], index_sheet_name: str = "01_Index") -> None:
    workbook = writer.book
    index_sheet = workbook[index_sheet_name]
    for row_index, sheet_name in enumerate(sheet_names, start=2):
        link_cell = index_sheet.cell(row=row_index, column=1)
        link_cell.value = sheet_name
        link_cell.hyperlink = _sheet_link(sheet_name, "A1")
        link_cell.style = "Hyperlink"

    for sheet_name in sheet_names:
        worksheet = workbook[sheet_name]
        worksheet["A1"] = "← Back to Index"
        worksheet["A1"].hyperlink = _sheet_link(index_sheet_name, "A1")
        worksheet["A1"].style = "Hyperlink"


def _append_excel_warning(warnings: list[dict[str, Any]], sheet_name: str, original_rows: int, written_rows: int) -> None:
    warnings.append({
        "Timestamp": datetime.now().isoformat(timespec="seconds"),
        "Level": "WARNING",
        "Source": "excel_report",
        "Message": "Excel sheet was truncated because it exceeded max_excel_rows_per_sheet.",
        "Context": {"sheet": sheet_name, "original_rows": original_rows, "written_rows": written_rows},
    })


def _truncate_frame_if_needed(
    sheet_name: str,
    frame: pd.DataFrame,
    max_rows: int,
    truncate_large_sheets: bool,
    warnings: list[dict[str, Any]],
) -> pd.DataFrame:
    original_rows = len(frame)
    safe_limit = min(max_rows, EXCEL_DATA_ROW_LIMIT)
    if original_rows <= safe_limit:
        return frame
    written_rows = safe_limit if truncate_large_sheets else EXCEL_DATA_ROW_LIMIT
    written_rows = min(written_rows, original_rows)
    _append_excel_warning(warnings, sheet_name, original_rows, written_rows)
    return frame.head(written_rows).copy()


# --- sheet-specific helpers (new) --------------------------------------------

SHEET_DESCRIPTIONS = {
    "02_Input": "Run configuration (flattened) and any warnings raised during this run.",
    "03_Theoretical": "Theoretical RNase digestion fragments and their unmodified monoisotopic mass.",
    "04_Observed_Mass": "Observed MS1 peaks with the charge(s) under which each peak matched a fragment in 06_Mass_Comparison.",
    "05_Mass_Intensity": "Same as 04_Observed_Mass with Intensity as an explicit column.",
    "06_Mass_Comparison": "Fragment x charge x Peak matches: mass differences plus independent Formula/Modification Candidate columns.",
    "07_Modifications": "Known RNA modification database (data/modifications.yaml) used for Modification Candidate search.",
    "07b_Hypothesis_Check": "Modification-hypothesis presence check: each hypothesis' theoretical mass (tRNA library defaults + config.hypothesis_check.targets) vs the raw peaks. Yes/No only — no automatic identification.",
    "08_Visualization": "Scatter chart: Charge (x) vs ΔDa (y), sourced from 06_Mass_Comparison and colored by whether a Formula/Modification candidate was found.",
    # §24 (Phase 10) P1 complete-digestion mode: 03/06 are repurposed (see
    # write_nucleoside_mass_hunter_report), 04/05/07 are unchanged (peaks
    # and the modification database don't depend on fragment vs nucleoside
    # mode), 08 sources from 06_Nucleoside_Comparison instead.
    "03_Nucleoside_Targets": "Known nucleoside mass universe (4 standard bases + data/modifications.yaml) used for P1 complete-digestion matching, under the current alkaline_phosphatase.enabled setting.",
    "06_Nucleoside_Comparison": "Peak x known-nucleoside-target matches (narrow ppm tolerance) for Nuclease P1 complete-digestion mode.",
}


def _modification_display_name(modification: Modification) -> str:
    raw = modification.raw or {}
    return str(raw.get("name") or modification.symbol or modification.id)


def _theoretical_frame(fragments: list[Fragment]) -> pd.DataFrame:
    rows = [
        {
            "Fragment ID": fragment.fragment_id,
            "Sequence": fragment.sequence,
            "Start": fragment.start,
            "End": fragment.end,
            "Length": fragment.end - fragment.start + 1,
            "Enzyme": fragment.enzyme,
            "Missed Cleavage": fragment.missed_cleavages,
            "Terminal Form": fragment.terminal_form,
            "Theoretical Mass": fragment.unmodified_mass,
        }
        for fragment in fragments
    ]
    return pd.DataFrame(rows, columns=["Fragment ID", "Sequence", "Start", "End", "Length", "Enzyme", "Missed Cleavage", "Terminal Form", "Theoretical Mass"])


def _format_rt_range(scan_count: int, rt_range: tuple[float, float] | None) -> str:
    """仕様書 §14C: scan_count<=1（cross-scanマージが起きていない通常の
    ピーク）なら空欄のまま。既存のRT列と重複する情報を増やさないため。"""
    if scan_count <= 1 or rt_range is None:
        return ""
    return f"{rt_range[0]:.3f}-{rt_range[1]:.3f}"


def _observed_rows(peaks: list[Peak], mass_comparison_rows: list[MassComparisonRow] | list[NucleosideComparisonRow]) -> list[dict[str, Any]]:
    """仕様書 §15/§16.2 の推奨に従い、1ピークにつき複数chargeでマッチした
    場合はcharge単位で行を分ける。マッチが1件も無かったピークも
    (charge/observed massは空欄のまま) 1行として残し、抽出済みMS1ピーク
    全体が04/05シートから欠落しないようにする。Scan Count/RT Range列は
    peak_picking.merge_peaks_across_scans() が統合したピークについてのみ
    埋まる（仕様書 §14C）。`mass_comparison_rows`はMassComparisonRow /
    NucleosideComparisonRowのどちらでもよい（peak_id/charge/observed_mass
    属性のみ参照するため、§24のP1モードでもそのまま再利用できる）。"""
    ids_by_index = assign_peak_ids(peaks)
    charges_by_peak_id: dict[str, set[int]] = {}
    observed_mass_by_peak_charge: dict[tuple[str, int], float] = {}
    for row in mass_comparison_rows:
        charges_by_peak_id.setdefault(row.peak_id, set()).add(row.charge)
        observed_mass_by_peak_charge.setdefault((row.peak_id, row.charge), row.observed_mass)

    rows: list[dict[str, Any]] = []
    for index, peak in enumerate(peaks):
        peak_id = ids_by_index[index]
        scan_count = getattr(peak, "scan_count", 1)
        rt_range_str = _format_rt_range(scan_count, getattr(peak, "rt_range", None))
        charges = sorted(charges_by_peak_id.get(peak_id, set()))
        if not charges:
            rows.append({
                "Peak ID": peak_id, "m/z": peak.mz, "Charge": None, "Observed Mass": None,
                "Intensity": peak.intensity, "RT": peak.rt, "Scan ID": peak.scan_id,
                "Scan Count": scan_count, "RT Range": rt_range_str,
            })
            continue
        for charge in charges:
            rows.append({
                "Peak ID": peak_id, "m/z": peak.mz, "Charge": charge,
                "Observed Mass": observed_mass_by_peak_charge.get((peak_id, charge)),
                "Intensity": peak.intensity, "RT": peak.rt, "Scan ID": peak.scan_id,
                "Scan Count": scan_count, "RT Range": rt_range_str,
            })
    return rows


def _observed_mass_frame(rows: list[dict[str, Any]]) -> pd.DataFrame:
    columns = ["Peak ID", "m/z", "Charge", "Observed Mass", "RT", "Scan ID", "Scan Count", "RT Range"]
    return pd.DataFrame(rows, columns=columns)


def _mass_intensity_frame(rows: list[dict[str, Any]]) -> pd.DataFrame:
    columns = ["Peak ID", "m/z", "Charge", "Observed Mass", "Intensity", "RT", "Scan ID", "Scan Count", "RT Range"]
    return pd.DataFrame(rows, columns=columns)


_MASS_COMPARISON_COLUMNS = [
    "Peak ID", "Fragment ID", "Sequence", "Charge", "Intensity", "Observed Mass", "Theoretical Mass",
    "ΔDa", "Δppm", "Recommended Formula", "Formula Candidates", "Known Modification",
    "Modification Candidates", "MS2 Spectrum ID", "MS2 Matched Ions", "Scan Count", "RT Range",
]


def _mass_comparison_frame(rows: list[MassComparisonRow]) -> pd.DataFrame:
    data = [
        {
            "Peak ID": row.peak_id, "Fragment ID": row.fragment_id, "Sequence": row.sequence,
            "Charge": row.charge, "Intensity": row.intensity, "Observed Mass": row.observed_mass,
            "Theoretical Mass": row.theoretical_mass, "ΔDa": row.delta_da, "Δppm": row.delta_ppm,
            "Recommended Formula": row.recommended_formula or "", "Formula Candidates": row.formula_candidates,
            "Known Modification": row.known_modification or "", "Modification Candidates": row.modification_candidates,
            "MS2 Spectrum ID": row.ms2_spectrum_id or "", "MS2 Matched Ions": row.ms2_matched_ions,
            "Scan Count": row.scan_count, "RT Range": _format_rt_range(row.scan_count, row.rt_range),
        }
        for row in rows
    ]
    return pd.DataFrame(data, columns=_MASS_COMPARISON_COLUMNS)


def _modifications_frame(modifications: list[Modification]) -> pd.DataFrame:
    rows = [
        {
            "ID": modification.id, "Symbol": modification.symbol or "", "Name": _modification_display_name(modification),
            "Target Bases": ", ".join(modification.target_bases), "Category": modification.category,
            "Mass Shift (Da)": modification.mass_shift_from_unmodified, "Chemical Group": modification.chemical_group,
            "Near-Isobaric Group": modification.near_isobaric_group, "Detectability": modification.detectability,
            "Curation Status": modification.curation_status,
        }
        for modification in modifications
    ]
    columns = ["ID", "Symbol", "Name", "Target Bases", "Category", "Mass Shift (Da)", "Chemical Group", "Near-Isobaric Group", "Detectability", "Curation Status"]
    return pd.DataFrame(rows, columns=columns)


def _target_universe_frame(targets: list[NucleosideTarget], dephosphorylated: bool, base_masses: dict[str, Any]) -> pd.DataFrame:
    """§24 (Phase 10) 03_Nucleoside_Targets: 標準4塩基+修飾ヌクレオシドの
    既知質量ユニバースを、現在の脱リン酸化設定下での理論質量とともに示す。"""
    rows = [
        {
            "Label": target.label, "Name": target.name, "Category": target.category,
            "Theoretical Mass": theoretical_mass_for_target(target, dephosphorylated, base_masses),
        }
        for target in targets
    ]
    return pd.DataFrame(rows, columns=["Label", "Name", "Category", "Theoretical Mass"])


_NUCLEOSIDE_COMPARISON_COLUMNS = [
    "Peak ID", "Target Label", "Target Name", "Category", "Charge", "Intensity",
    "Observed Mass", "Theoretical Mass", "ΔDa", "Δppm", "Scan Count", "RT Range",
]


def _nucleoside_comparison_frame(rows: list[NucleosideComparisonRow]) -> pd.DataFrame:
    """§24 (Phase 10) 06_Nucleoside_Comparison: 06_Mass_Comparisonと同じ
    Index/ハイパーリンク/切り詰め/書式の枠組みを流用するが、列はFragment
    ではなくNucleosideTargetベース（Recommended Formula等のFormula
    Candidate関連列は無い——P1モードは既知質量照合のみがスコープ、§24）。"""
    data = [
        {
            "Peak ID": row.peak_id, "Target Label": row.target_label, "Target Name": row.target_name,
            "Category": row.target_category, "Charge": row.charge, "Intensity": row.intensity,
            "Observed Mass": row.observed_mass, "Theoretical Mass": row.theoretical_mass,
            "ΔDa": row.delta_da, "Δppm": row.delta_ppm,
            "Scan Count": row.scan_count, "RT Range": _format_rt_range(row.scan_count, row.rt_range),
        }
        for row in rows
    ]
    return pd.DataFrame(data, columns=_NUCLEOSIDE_COMPARISON_COLUMNS)


_HYPOTHESIS_CHECK_COLUMNS = [
    "Hypothesis_Name", "Source", "Formula_Description", "Theoretical_Mass", "Match_Found",
    "Matched_Peak_IDs", "Charge", "Observed_Mass", "ΔDa", "Δppm", "Intensity", "RT_Range",
]


def _hypothesis_rt_range(row: HypothesisCheckRow) -> str:
    """Merged peaks (§14C) show their RT span; an ordinary peak shows its
    single RT — for a hypothesis check the elution time is useful even when
    no merging happened (unlike the other sheets' RT Range column)."""
    if row.scan_count > 1 and row.rt_range is not None:
        return _format_rt_range(row.scan_count, row.rt_range)
    return "" if row.rt is None else f"{row.rt:.3f}"


def _hypothesis_check_frame(rows: list[HypothesisCheckRow]) -> pd.DataFrame:
    """hypothesis_mass_check_spec.md §6: one row per matching (peak, charge);
    a hypothesis with no match gets a single Match_Found=No row."""
    data = [
        {
            "Hypothesis_Name": row.hypothesis_name, "Source": row.source,
            "Formula_Description": row.formula_description, "Theoretical_Mass": row.theoretical_mass,
            "Match_Found": "Yes" if row.match_found else "No", "Matched_Peak_IDs": row.peak_id or "",
            "Charge": row.charge, "Observed_Mass": row.observed_mass, "ΔDa": row.delta_da,
            "Δppm": row.delta_ppm, "Intensity": row.intensity, "RT_Range": _hypothesis_rt_range(row),
        }
        for row in rows
    ]
    return pd.DataFrame(data, columns=_HYPOTHESIS_CHECK_COLUMNS)


def _input_frame(config: RunConfig, warnings: list[dict[str, Any]]) -> pd.DataFrame:
    config_dict = {key: value for key, value in vars(config).items() if key != "raw"}
    rows = _flatten_dict(config_dict)
    if warnings:
        rows.append({"Parameter": "--- Warnings ---", "Value": f"{len(warnings)} warning(s)"})
        for warning in warnings:
            context = warning.get("Context")
            suffix = f" ({context!r})" if context not in (None, "") else ""
            rows.append({"Parameter": f"[{warning['Level']}] {warning['Source']}", "Value": f"{warning['Message']}{suffix}"})
    return pd.DataFrame(rows, columns=["Parameter", "Value"])


_VIZ_HELPER_COL = 20  # column T — far enough right to stay clear of the chart anchored at A5


def _add_visualization_sheet(
    writer: pd.ExcelWriter,
    mass_comparison_rows: list[MassComparisonRow],
    config: RunConfig,
    warnings: list[dict[str, Any]],
) -> None:
    """仕様書 §17（2026-09-03改訂）: 08_Visualization に Charge(x) x ΔDa(y)
    の散布図を1枚配置する。データソースは04_Observed_Massではなく
    06_Mass_Comparison（各行がFormula/Modification候補を持つかどうかで
    色分けし、有力な候補が付いた質量差がchargeごとにどう分布しているかを
    一目で見られるようにする）。

    候補有無での色分けは、openpyxlのScatterChartが1系列内で点ごとに色を
    変える機能を持たないため、2系列（候補あり/候補なし）に分けて描画する
    必要がある。06シート本体の行順（fragment→charge順）を変えずに済む
    よう、行を並べ替えるのではなく、このシートの遠い列（T列以降）に
    候補あり/なし別のCharge・ΔDaを転記した小さな補助表を用意し、
    そちらをチャートのデータソースにする（値はopenpyxl経由でセルに直接
    書き込むだけで、06シートへのセル参照ではない——候補有無は集計値で
    あり単純なセル参照では表現できないため）。

    チャート生成が失敗しても例外を握りつぶし、レポート全体の出力は継続
    する（§17, §18）。"""
    sheet = writer.book.create_sheet("08_Visualization")
    sheet["A1"] = "← Back to Index"
    sheet["A1"].hyperlink = _sheet_link("01_Index", "A1")
    sheet["A1"].style = "Hyperlink"

    viz_config = config.visualization or {}
    if not viz_config.get("enabled", True):
        sheet["A3"] = "Visualization is disabled (config.visualization.enabled = false)."
        return
    if not mass_comparison_rows:
        sheet["A3"] = "No Mass Comparison rows to plot."
        return

    try:
        from openpyxl.chart import Reference, ScatterChart, Series
        from openpyxl.chart.marker import Marker
        from openpyxl.chart.shapes import GraphicalProperties

        with_candidate = [row for row in mass_comparison_rows if has_any_candidate(row)]
        without_candidate = [row for row in mass_comparison_rows if not has_any_candidate(row)]

        header_row = 1
        col = _VIZ_HELPER_COL
        sheet.cell(row=header_row, column=col, value="Charge (has candidate)")
        sheet.cell(row=header_row, column=col + 1, value="ΔDa (has candidate)")
        sheet.cell(row=header_row, column=col + 2, value="Charge (no candidate)")
        sheet.cell(row=header_row, column=col + 3, value="ΔDa (no candidate)")
        for offset, row in enumerate(with_candidate, start=header_row + 1):
            sheet.cell(row=offset, column=col, value=row.charge)
            sheet.cell(row=offset, column=col + 1, value=row.delta_da)
        for offset, row in enumerate(without_candidate, start=header_row + 1):
            sheet.cell(row=offset, column=col + 2, value=row.charge)
            sheet.cell(row=offset, column=col + 3, value=row.delta_da)

        chart = ScatterChart()
        chart.title = "ΔDa by Charge (colored by candidate presence)"
        chart.x_axis.title = "Charge"
        chart.y_axis.title = "ΔDa (Da)"
        chart.style = 2

        def _series(mz_col: int, first_row: int, last_row: int, title: str, color: str) -> Series:
            x_values = Reference(sheet, min_col=col + mz_col, min_row=first_row, max_row=last_row)
            y_values = Reference(sheet, min_col=col + mz_col + 1, min_row=first_row, max_row=last_row)
            series = Series(y_values, x_values, title=title)
            series.marker = Marker(symbol="circle")
            series.marker.graphicalProperties = GraphicalProperties(solidFill=color)
            series.graphicalProperties.line.noFill = True
            return series

        if with_candidate:
            chart.series.append(_series(0, header_row + 1, header_row + len(with_candidate), "Has candidate", "2E75B6"))
        if without_candidate:
            chart.series.append(_series(2, header_row + 1, header_row + len(without_candidate), "No candidate", "BFBFBF"))

        sheet.add_chart(chart, "A5")
    except Exception as exc:  # noqa: BLE001 - chart generation must never abort the whole report
        from rna_masshunter.warnings_manager import add_warning

        add_warning(warnings, "WARNING", "excel_report", "08_Visualization chart generation failed; report was written without it.", str(exc))
        sheet["A3"] = "Chart generation failed; see Warnings in 02_Input."


def _add_nucleoside_visualization_sheet(
    writer: pd.ExcelWriter,
    nucleoside_comparison_rows: list[NucleosideComparisonRow],
    config: RunConfig,
    warnings: list[dict[str, Any]],
) -> None:
    """§24 (Phase 10) P1モード版の08_Visualization。
    _add_visualization_sheet と同じ構造だが、候補有無ではなく
    target_category（standard/modified）で2系列に色分けする。"""
    sheet = writer.book.create_sheet("08_Visualization")
    sheet["A1"] = "← Back to Index"
    sheet["A1"].hyperlink = _sheet_link("01_Index", "A1")
    sheet["A1"].style = "Hyperlink"

    viz_config = config.visualization or {}
    if not viz_config.get("enabled", True):
        sheet["A3"] = "Visualization is disabled (config.visualization.enabled = false)."
        return
    if not nucleoside_comparison_rows:
        sheet["A3"] = "No Nucleoside Comparison rows to plot."
        return

    try:
        from openpyxl.chart import Reference, ScatterChart, Series
        from openpyxl.chart.marker import Marker
        from openpyxl.chart.shapes import GraphicalProperties

        standard_rows = [row for row in nucleoside_comparison_rows if row.target_category == "standard"]
        modified_rows = [row for row in nucleoside_comparison_rows if row.target_category != "standard"]

        header_row = 1
        col = _VIZ_HELPER_COL
        sheet.cell(row=header_row, column=col, value="Charge (standard)")
        sheet.cell(row=header_row, column=col + 1, value="ΔDa (standard)")
        sheet.cell(row=header_row, column=col + 2, value="Charge (modified)")
        sheet.cell(row=header_row, column=col + 3, value="ΔDa (modified)")
        for offset, row in enumerate(standard_rows, start=header_row + 1):
            sheet.cell(row=offset, column=col, value=row.charge)
            sheet.cell(row=offset, column=col + 1, value=row.delta_da)
        for offset, row in enumerate(modified_rows, start=header_row + 1):
            sheet.cell(row=offset, column=col + 2, value=row.charge)
            sheet.cell(row=offset, column=col + 3, value=row.delta_da)

        chart = ScatterChart()
        chart.title = "ΔDa by Charge (colored by standard vs modified nucleoside)"
        chart.x_axis.title = "Charge"
        chart.y_axis.title = "ΔDa (Da)"
        chart.style = 2

        def _series(mz_col: int, first_row: int, last_row: int, title: str, color: str) -> Series:
            x_values = Reference(sheet, min_col=col + mz_col, min_row=first_row, max_row=last_row)
            y_values = Reference(sheet, min_col=col + mz_col + 1, min_row=first_row, max_row=last_row)
            series = Series(y_values, x_values, title=title)
            series.marker = Marker(symbol="circle")
            series.marker.graphicalProperties = GraphicalProperties(solidFill=color)
            series.graphicalProperties.line.noFill = True
            return series

        if standard_rows:
            chart.series.append(_series(0, header_row + 1, header_row + len(standard_rows), "Standard base", "2E75B6"))
        if modified_rows:
            chart.series.append(_series(2, header_row + 1, header_row + len(modified_rows), "Modified nucleoside", "C0504D"))

        sheet.add_chart(chart, "A5")
    except Exception as exc:  # noqa: BLE001 - chart generation must never abort the whole report
        from rna_masshunter.warnings_manager import add_warning

        add_warning(warnings, "WARNING", "excel_report", "08_Visualization chart generation failed; report was written without it.", str(exc))
        sheet["A3"] = "Chart generation failed; see Warnings in 02_Input."


def _apply_number_formats(worksheet: Any, frame: pd.DataFrame, formats: dict[str, str], header_row: int = 3) -> None:
    if frame.empty:
        return
    for column_name, number_format in formats.items():
        if column_name not in frame.columns:
            continue
        column_index = list(frame.columns).index(column_name) + 1
        column_letter = get_column_letter(column_index)
        for row_offset in range(len(frame)):
            cell = worksheet[f"{column_letter}{header_row + 1 + row_offset}"]
            cell.number_format = number_format


def write_simple_mass_hunter_report(
    output_path: str | Path,
    config: RunConfig,
    fragments: list[Fragment],
    peaks: list[Peak],
    mass_comparison_rows: list[MassComparisonRow],
    modifications: list[Modification],
    warnings: list[dict[str, Any]] | None = None,
    hypothesis_rows: list[HypothesisCheckRow] | None = None,
) -> Path:
    """仕様書 §16: 01_Index〜08_Visualization を出力する。既存の
    _add_index_and_backlinks/_autosize_and_freeze のIndex/ハイパーリンク/
    書式の流儀をそのまま踏襲する（§16.3）。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    warnings = warnings if warnings is not None else []

    reporting = config.reporting or {}
    max_rows = int(reporting.get("max_excel_rows_per_sheet", 100000) or 100000)
    truncate_large_sheets = bool(reporting.get("truncate_large_sheets", True))

    observed_rows = _observed_rows(peaks, mass_comparison_rows)

    sheets: dict[str, pd.DataFrame] = {
        "02_Input": _input_frame(config, warnings),
        "03_Theoretical": _theoretical_frame(fragments),
        "04_Observed_Mass": _observed_mass_frame(observed_rows),
        "05_Mass_Intensity": _mass_intensity_frame(observed_rows),
        "06_Mass_Comparison": _mass_comparison_frame(mass_comparison_rows),
        "07_Modifications": _modifications_frame(modifications),
    }
    if hypothesis_rows is not None:
        sheets["07b_Hypothesis_Check"] = _hypothesis_check_frame(hypothesis_rows)
    sheets = {name: _truncate_frame_if_needed(name, frame, max_rows, truncate_large_sheets, warnings) for name, frame in sheets.items()}
    # Cells may hold None / dict (e.g. Detectability) — normalize them so
    # openpyxl never chokes on an unsupported type (仕様書 §18 error handling
    # spirit: prefer a readable blank/JSON string over a crash).
    sheets = {name: frame.map(_excel_safe_cell) if not frame.empty else frame for name, frame in sheets.items()}

    all_sheet_names = list(sheets) + ["08_Visualization"]
    index_rows = [{"Sheet": name, "Description": SHEET_DESCRIPTIONS.get(name, ""), "Notes": "Data starts at A3."} for name in all_sheet_names]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame(index_rows, columns=["Sheet", "Description", "Notes"]).to_excel(writer, sheet_name="01_Index", index=False)
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False, startrow=2)
        _add_visualization_sheet(writer, mass_comparison_rows, config, warnings)
        _add_index_and_backlinks(writer, all_sheet_names, index_sheet_name="01_Index")

        # 仕様書 §16.4: mass系列は"0.000"、ΔDa系列は符号付き"+0.00000;-0.00000"、Δppmは"0.0"。
        _apply_number_formats(writer.sheets["03_Theoretical"], sheets["03_Theoretical"], {"Theoretical Mass": "0.000"})
        _apply_number_formats(writer.sheets["04_Observed_Mass"], sheets["04_Observed_Mass"], {"m/z": "0.000", "Observed Mass": "0.000"})
        _apply_number_formats(writer.sheets["05_Mass_Intensity"], sheets["05_Mass_Intensity"], {"m/z": "0.000", "Observed Mass": "0.000"})
        _apply_number_formats(writer.sheets["06_Mass_Comparison"], sheets["06_Mass_Comparison"], {
            "Observed Mass": "0.000", "Theoretical Mass": "0.000",
            "ΔDa": "+0.00000;-0.00000", "Δppm": "0.0",
        })

        if "07b_Hypothesis_Check" in sheets:
            _apply_number_formats(writer.sheets["07b_Hypothesis_Check"], sheets["07b_Hypothesis_Check"], {
                "Theoretical_Mass": "0.0000", "Observed_Mass": "0.0000",
                "ΔDa": "+0.00000;-0.00000", "Δppm": "0.0",
            })

        _autosize_and_freeze(writer, index_sheet_name="01_Index")

    return output_path


def write_nucleoside_mass_hunter_report(
    output_path: str | Path,
    config: RunConfig,
    targets: list[NucleosideTarget],
    peaks: list[Peak],
    nucleoside_comparison_rows: list[NucleosideComparisonRow],
    modifications: list[Modification],
    base_masses: dict[str, Any],
    warnings: list[dict[str, Any]] | None = None,
    hypothesis_rows: list[HypothesisCheckRow] | None = None,
) -> Path:
    """§24 (Phase 10): Nuclease P1完全分解モード用のExcel出力。
    write_simple_mass_hunter_report と同じ01_Index〜08_Visualizationの
    枠組み（Index/ハイパーリンク/切り詰め/書式まわりの共通ヘルパー）を
    そのまま流用するが、03/06シートはFragmentベースではなくNucleosideTarget
    ベースの内容に差し替える（04/05/07は共通——ピークの生データと
    修飾データベースはどちらのモードでも同じ）。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    warnings = warnings if warnings is not None else []

    reporting = config.reporting or {}
    max_rows = int(reporting.get("max_excel_rows_per_sheet", 100000) or 100000)
    truncate_large_sheets = bool(reporting.get("truncate_large_sheets", True))
    dephosphorylated = bool((config.alkaline_phosphatase or {}).get("enabled", False))

    observed_rows = _observed_rows(peaks, nucleoside_comparison_rows)

    sheets: dict[str, pd.DataFrame] = {
        "02_Input": _input_frame(config, warnings),
        "03_Nucleoside_Targets": _target_universe_frame(targets, dephosphorylated, base_masses),
        "04_Observed_Mass": _observed_mass_frame(observed_rows),
        "05_Mass_Intensity": _mass_intensity_frame(observed_rows),
        "06_Nucleoside_Comparison": _nucleoside_comparison_frame(nucleoside_comparison_rows),
        "07_Modifications": _modifications_frame(modifications),
    }
    if hypothesis_rows is not None:
        sheets["07b_Hypothesis_Check"] = _hypothesis_check_frame(hypothesis_rows)
    sheets = {name: _truncate_frame_if_needed(name, frame, max_rows, truncate_large_sheets, warnings) for name, frame in sheets.items()}
    sheets = {name: frame.map(_excel_safe_cell) if not frame.empty else frame for name, frame in sheets.items()}

    all_sheet_names = list(sheets) + ["08_Visualization"]
    index_rows = [{"Sheet": name, "Description": SHEET_DESCRIPTIONS.get(name, ""), "Notes": "Data starts at A3."} for name in all_sheet_names]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame(index_rows, columns=["Sheet", "Description", "Notes"]).to_excel(writer, sheet_name="01_Index", index=False)
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False, startrow=2)
        _add_nucleoside_visualization_sheet(writer, nucleoside_comparison_rows, config, warnings)
        _add_index_and_backlinks(writer, all_sheet_names, index_sheet_name="01_Index")

        _apply_number_formats(writer.sheets["03_Nucleoside_Targets"], sheets["03_Nucleoside_Targets"], {"Theoretical Mass": "0.000"})
        _apply_number_formats(writer.sheets["04_Observed_Mass"], sheets["04_Observed_Mass"], {"m/z": "0.000", "Observed Mass": "0.000"})
        _apply_number_formats(writer.sheets["05_Mass_Intensity"], sheets["05_Mass_Intensity"], {"m/z": "0.000", "Observed Mass": "0.000"})
        _apply_number_formats(writer.sheets["06_Nucleoside_Comparison"], sheets["06_Nucleoside_Comparison"], {
            "Observed Mass": "0.000", "Theoretical Mass": "0.000",
            "ΔDa": "+0.00000;-0.00000", "Δppm": "0.0",
        })

        if "07b_Hypothesis_Check" in sheets:
            _apply_number_formats(writer.sheets["07b_Hypothesis_Check"], sheets["07b_Hypothesis_Check"], {
                "Theoretical_Mass": "0.0000", "Observed_Mass": "0.0000",
                "ΔDa": "+0.00000;-0.00000", "Δppm": "0.0",
            })

        _autosize_and_freeze(writer, index_sheet_name="01_Index")

    return output_path
